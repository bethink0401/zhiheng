from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT_DIR = Path(__file__).resolve().parent
OUT = OUT_DIR / "知衡国创赛项目计划书.docx"
LOGO = OUT_DIR.parent / "Design" / "AppIcon-Zhiheng-v3-focus-liquid-glass.png"

BLUE = "176B87"
DEEP = "103B4D"
TEAL = "15A7A0"
LIGHT = "EAF6F5"
PALE = "F4F8FA"
GRAY = "60707A"
LINE = "D8E3E7"
WHITE = "FFFFFF"


def set_fonts(run, size=11, bold=False, color=None, font="Hiragino Sans GB"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run._element.rPr.rFonts.set(qn("w:cs"), font)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margin(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for old in list(grid):
        grid.remove(old)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margin(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color=LINE, sz="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:color"), color)


def add_text(cell, text, size=10.5, bold=False, color=DEEP, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.18
    r = p.add_run(text)
    set_fonts(r, size=size, bold=bold, color=color)
    return p


def add_para(doc, text="", size=11, bold=False, color="1D2930", align=None, before=0, after=8, line=1.333):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    r = p.add_run(text)
    set_fonts(r, size=size, bold=bold, color=color)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.194)
    r = p.add_run(text)
    set_fonts(r, size=10.5, color="1D2930")
    return p


def add_h1(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run(text)
    set_fonts(r, size=16, bold=True, color=BLUE)
    return p


def add_h2(doc, text):
    p = doc.add_paragraph(style="Heading 2")
    r = p.add_run(text)
    set_fonts(r, size=13, bold=True, color=BLUE)
    return p


def add_h3(doc, text):
    p = doc.add_paragraph(style="Heading 3")
    r = p.add_run(text)
    set_fonts(r, size=12, bold=True, color=DEEP)
    return p


def add_callout(doc, title, body, color=TEAL):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color=color, sz="10")
    cell = table.cell(0, 0)
    shade(cell, LIGHT)
    set_cell_margin(cell, top=150, start=220, bottom=150, end=220)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_fonts(r, size=10.5, bold=True, color=DEEP)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.25
    r2 = p2.add_run(body)
    set_fonts(r2, size=10.5, color=DEEP)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    for cell, text in zip(table.rows[0].cells, headers):
        shade(cell, DEEP)
        add_text(cell, text, size=10, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = table.add_row().cells
        for i, (cell, text) in enumerate(zip(cells, row)):
            if len(table.rows) % 2 == 1:
                shade(cell, PALE)
            add_text(cell, text, size=9.7, color="20303A", align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_flow(doc, labels):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    for i, label in enumerate(labels):
        r = p.add_run(label)
        set_fonts(r, size=10, bold=True, color=DEEP)
        if i != len(labels) - 1:
            arrow = p.add_run("  →  ")
            set_fonts(arrow, size=12, bold=True, color=TEAL)


def add_page_break(doc):
    doc.add_page_break()


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Hiragino Sans GB"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Hiragino Sans GB")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Hiragino Sans GB")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Hiragino Sans GB")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DEEP, 8, 4),
    ]:
        st = styles[style_name]
        st.font.name = "Hiragino Sans GB"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
        st._element.rPr.rFonts.set(qn("w:ascii"), "Hiragino Sans GB")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Hiragino Sans GB")
        st._element.rPr.rFonts.set(qn("w:cs"), "Hiragino Sans GB")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("知衡  |  国创赛项目计划书")
    set_fonts(hr, size=8.5, color=GRAY)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("健康数据用于理解与行动，不用于诊断或治疗")
    set_fonts(fr, size=8.5, color=GRAY)


def cover(doc):
    doc.add_paragraph().paragraph_format.space_after = Pt(55)
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Inches(1.18))
    add_para(doc, "国创赛项目计划书", size=12, bold=True, color=TEAL, align=WD_ALIGN_PARAGRAPH.CENTER, before=18, after=12)
    add_para(doc, "知衡", size=30, bold=True, color=DEEP, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_para(doc, "从健康洞察到行为改变的AI个人健康智能体", size=17, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=20)
    add_para(doc, "让健康数据转化为每个人可理解、可执行、可验证的有效方法", size=11.5, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=52)
    t = doc.add_table(rows=4, cols=2)
    set_table_geometry(t, [1900, 7460])
    set_table_borders(t)
    data = [
        ("项目名称", "知衡——从健康洞察到行为改变的AI个人健康智能体"),
        ("参赛赛道", "高教主赛道（请按学校最终通知填写）"),
        ("申报单位", "【待填写：学校名称】"),
        ("项目团队", "【待填写：负责人、成员及指导教师】"),
    ]
    for row, (label, value) in zip(t.rows, data):
        shade(row.cells[0], PALE)
        add_text(row.cells[0], label, size=10.5, bold=True, color=DEEP)
        add_text(row.cells[1], value, size=10.5, color="20303A")
    add_para(doc, "版本：V1.0  |  日期：2026年9月1日", size=9.5, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, before=40, after=0)
    add_page_break(doc)


def toc(doc):
    add_para(doc, "目录", size=24, bold=True, color=DEEP, after=20)
    entries = [
        "01 项目摘要", "02 行业背景与用户痛点", "03 产品方案与使用场景", "04 核心创新与竞争优势",
        "05 技术路线与AI安全", "06 市场定位与商业模式", "07 项目进展与实施计划", "08 团队建设与资源需求",
        "09 社会价值、风险与合规", "10 附录：待补材料清单",
    ]
    for entry in entries:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(entry)
        set_fonts(r, size=12, color=DEEP)
    add_callout(doc, "使用说明", "本稿已使用现有产品事实与项目定位撰写。方括号中的学校、团队、调研、合作、知识产权及财务数据须由团队补充后再提交；未完成的事项均以计划表述，不应替换为既有成果。")
    add_page_break(doc)


def body(doc):
    add_h1(doc, "01  项目摘要")
    add_para(doc, "知衡是一款从健康洞察走向行为改变的AI个人健康智能体。面对健康数据持续增长而用户仍然看不懂、用不好、坚持不下去的痛点，知衡将多维健康数据与用户的精力、压力、身体感受和生活情境连接起来，通过数据质量判断、个人基线建模、趋势识别和AI解释，帮助用户理解近期值得关注的变化。")
    add_para(doc, "知衡不止记录数据，而是构建“发现变化、理解变化、采取行动、验证效果、沉淀方法”的完整闭环。系统在每次洞察后只推荐一个三至七天、低风险、可随时停止的微计划；计划结束后，结合完成率、主观反馈、客观趋势和数据质量，形成“可能有帮助”“暂未观察到明显变化”“执行不足”或“数据不足”等审慎结论，逐步沉淀用户专属的有效方法库。")
    add_callout(doc, "项目使命", "让健康数据不再只是图表中的数字，而成为帮助用户理解自己、改善生活方式、持续积累个人健康方法的能力。")
    add_h2(doc, "项目价值主张")
    add_flow(doc, ["可信数据", "个体洞察", "微行动", "效果验证", "个人方法资产"])
    add_table(doc, ["用户得到什么", "知衡如何实现"], [
        ("看懂近期变化", "用个人基线与趋势识别替代单纯指标展示"),
        ("获得低负担行动", "一次只给出一个三至七天、低风险的微计划"),
        ("知道是否适合自己", "综合完成率、主观感受、趋势与数据质量进行评估"),
        ("长期形成健康经验", "将多次验证后的结果沉淀为个人有效方法库"),
    ], [3000, 6360])

    add_h1(doc, "02  行业背景与用户痛点")
    add_h2(doc, "2.1 数据增长与健康行动之间的断层")
    add_para(doc, "智能穿戴设备和健康数据平台正在让运动、睡眠、心率等信息更容易被记录，但“记录更多”并不等于“理解更多”。用户常常面对分散的图表、统一的目标和短期的打卡任务，仍难以回答自己最关心的三个问题：最近发生了什么，我可以如何改善，以及这种方法是否真的适合我。")
    add_h2(doc, "2.2 四类关键痛点")
    for item in [
        "数据碎片化：指标分散在不同页面，用户难以形成完整的自我理解。",
        "群体标准化：统一阈值忽视不同作息、体能和生活方式造成的个体差异。",
        "建议同质化：泛化的早睡、运动、少咖啡建议难以转化为持续行动。",
        "缺少验证闭环：用户完成行动后，通常无法判断是偶然波动还是值得保留的方法。",
    ]:
        add_bullet(doc, item)
    add_callout(doc, "核心问题", "如何把长期健康数据转化为可信、低焦虑、可执行、可验证的个人健康行动。", color=BLUE)
    add_h2(doc, "2.3 目标用户与优先场景")
    add_table(doc, ["优先用户", "典型场景", "知衡提供的价值"], [
        ("大学生", "考试、竞赛、作息波动", "帮助识别睡眠与精力变化，形成轻量改善方案"),
        ("青年职场人", "久坐、加班、高压力", "将日常状态和生活情境转化为可坚持的微行动"),
        ("运动入门用户", "训练恢复、习惯建立", "避免单日数据误导，支持观察与低风险调整"),
        ("关注家庭健康的成年人", "共同关注作息与生活方式", "沉淀长期、个体化的健康经验"),
    ], [1800, 3000, 4560])

    add_h1(doc, "03  产品方案与使用场景")
    add_h2(doc, "3.1 产品闭环")
    add_flow(doc, ["授权连接", "数据质量", "个人基线", "AI洞察", "微计划", "效果评估", "有效方法"])
    add_para(doc, "知衡在分析前先判断数据是否足以支持结论。对于数据缺失、来源变化、异常波动或设备未使用等情况，系统优先提示“数据不足”或“建议继续观察”，而不是强行输出结论。只有当近期有效数据与个人基线共同满足条件时，才生成值得继续观察的变化。")
    add_h2(doc, "3.2 核心功能")
    add_table(doc, ["功能模块", "用户体验", "差异化价值"], [
        ("健康数据总览", "查看近期状态与关键记录", "不以单一健康分数定义用户"),
        ("个人基线与趋势", "理解自己与过去的自己相比发生了什么", "以长期有效数据建立动态个人参考"),
        ("主观与情境记录", "快速记录精力、压力、身体感受和生活事件", "让客观数据与真实感受共同参与分析"),
        ("AI健康洞察", "获得事实、可能因素、不确定性与下一步建议", "AI只解释结构化健康事实，不自由编造数据"),
        ("微计划", "执行一个低风险、可停止的三至七天行动", "降低复杂计划与打卡挫败感"),
        ("我的有效方法", "回顾什么情境下什么方法可能有帮助", "把一次性建议转化为长期个人资产"),
    ], [1900, 3480, 3980])
    add_h2(doc, "3.3 典型用户故事")
    add_para(doc, "一名准备比赛的学生发现自己连续数日睡眠时长相较于个人近期基线出现变化，同时记录了较高压力和晚间咖啡情境。知衡不会断言原因，而是说明已观察到的事实、提示可能相关因素与其他可能性，并建议尝试一个“下午设定咖啡截止时间”的短期微计划。计划结束后，系统结合执行情况、主观精力与可用数据趋势给出审慎的本地评估。用户由此获得的不是抽象建议，而是一条可以继续验证的个人经验。")

    add_h1(doc, "04  核心创新与竞争优势")
    add_h2(doc, "4.1 五项核心创新")
    innovations = [
        ("个人基线引擎", "主要比较用户与过去的自己，降低统一标准对个体差异的误判。"),
        ("数据质量守门", "先识别有效天数、缺失、异常和来源变化，再决定是否输出洞察。"),
        ("主客观融合", "把用户感受和生活情境与客观数据并列，不让设备数据否定人的体验。"),
        ("AI解释与微计划闭环", "AI在事实边界内解释趋势并引导一个低风险行动，而不是泛化说教。"),
        ("个人有效方法资产", "通过多次行动与评估，沉淀可回看、可复用、可继续验证的方法库。"),
    ]
    add_table(doc, ["创新点", "创新说明"], innovations, [2680, 6680])
    add_h2(doc, "4.2 与常见产品类别的差异")
    add_table(doc, ["对比维度", "数据展示类应用", "通用AI问答", "知衡"], [
        ("健康数据处理", "以图表和单项指标为主", "依赖用户描述", "数据质量与个人基线先行"),
        ("健康解释", "有限或固定", "可能脱离真实数据", "基于结构化事实包与时间范围"),
        ("行动建议", "统一目标或打卡", "临时建议", "一次一个可停止的微计划"),
        ("效果验证", "通常缺失", "通常缺失", "完成率、感受、趋势和数据质量共同评估"),
        ("长期价值", "历史图表", "一次性对话", "个人有效方法库"),
    ], [1900, 2480, 2480, 2500])
    add_callout(doc, "竞争壁垒", "知衡的竞争力不在于“增加一个AI聊天入口”，而在于把个人基线、数据质量、生活情境、行动记录和效果验证连接为一个可追溯的长期系统。")

    add_h1(doc, "05  技术路线与AI安全")
    add_h2(doc, "5.1 技术路线")
    add_flow(doc, ["数据接入", "本地清洗", "趋势计算", "健康事实包", "安全规则", "AI解释", "客户端校验"])
    add_para(doc, "知衡采用“本地确定性分析加受控AI解释”的技术路线。数据质量、个人基线、趋势结果和计划评估由可测试的本地规则生成；AI只使用完成回答所必需的聚合事实，用于解释、追问和行动引导。这样既发挥AI的自然语言交互能力，也保留数据依据和可追溯性。")
    add_h2(doc, "5.2 数据与AI边界")
    add_table(doc, ["环节", "设计原则", "用户价值"], [
        ("数据处理", "本地优先，只使用实现当前功能所需的数据", "减少敏感数据暴露"),
        ("事实包", "仅发送时间范围、聚合趋势、数据质量与用户主动提供的相关信息", "AI回答有明确依据"),
        ("AI输出", "区分事实、可能因素、不确定性、一个追问和建议行动", "避免把相关性写成因果"),
        ("安全规则", "识别诊断、药物和紧急症状请求并优先安全处理", "明确非诊疗边界"),
        ("失败降级", "网络或模型失败时保留本地规则摘要", "核心体验不依赖AI在线可用"),
    ], [1760, 4550, 3050])
    add_h2(doc, "5.3 安全与隐私承诺")
    for item in [
        "不提供疾病诊断、处方、药物剂量调整或持续急救监护。",
        "不将单项指标直接等同于压力、恢复能力或疾病。",
        "不把健康数据用于广告、出售画像或训练基础模型。",
        "原始健康数据默认不上传；用户可查看权限状态并删除相关记录。",
        "当用户主动描述紧急症状时，优先提示及时寻求当地急救或专业帮助。",
    ]:
        add_bullet(doc, item)

    add_h1(doc, "06  市场定位与商业模式")
    add_h2(doc, "6.1 市场定位")
    add_para(doc, "知衡定位于个人健康管理和生活方式改善，而非诊疗服务。项目以成年个人用户为起点，先在高频、高压力、对数据理解有明确需求的学习与工作场景中验证价值，再逐步探索家庭、校园、企业与社区健康促进场景。")
    add_h2(doc, "6.2 阶段化服务路径")
    add_table(doc, ["阶段", "重点用户与场景", "核心目标"], [
        ("第一阶段", "大学生与青年职场人", "验证健康洞察、微计划和有效方法闭环"),
        ("第二阶段", "睡眠、运动恢复、久坐和压力管理场景", "扩展可验证的健康行为改善模板"),
        ("第三阶段", "家庭、校园、企业和社区", "探索群体健康促进与服务协同"),
        ("长期", "多设备与多健康数据平台", "形成开放、可信的个人健康服务网络"),
    ], [1600, 3860, 3900])
    add_h2(doc, "6.3 商业模式设想")
    add_para(doc, "项目当前以产品验证和用户价值验证为优先。长期可采用基础功能免费、进阶分析与个性化服务增值的模式，并探索面向校园和企业的健康促进服务。任何商业化均以用户授权、隐私保护和不影响健康建议独立性为前提，不以出售健康数据或健康广告为核心收入来源。")
    add_callout(doc, "提交提醒", "若学校要求填写营收、融资或市场规模，必须使用团队可提供的可核验材料和明确出处；本计划书不虚构订单、用户量、合作协议或收入。", color=BLUE)

    add_h1(doc, "07  项目进展与实施计划")
    add_h2(doc, "7.1 当前开发基础")
    add_para(doc, "截至2026年9月1日，项目已建立移动端产品工程、健康数据读取与演示数据适配、个人基线与趋势、主观记录、可解释洞察、AI流式对话、微计划执行、计划开始基线快照和结束评估等核心能力。项目已完成模拟器交互与构建验证，并通过162项iOS测试和14项AI代理测试。上述数字代表当前工程测试结果，不等同于临床验证或大规模用户验证。")
    add_table(doc, ["模块", "已完成的核心能力", "下一步重点"], [
        ("数据与洞察", "多类健康数据映射、数据质量、个人基线、7天与28天趋势", "继续完成真实数据与多来源验证"),
        ("AI能力", "结构化事实包、流式对话、输出校验和安全边界", "完善断网与真实环境验证"),
        ("行动闭环", "十类低风险模板、计划执行、反馈、暂停恢复、结束评估", "补齐计划更新、跨天、删除与历史一致性"),
        ("产品体验", "今日、洞察、微计划、AI助手、我的五个主要页面", "开展可理解性与低焦虑体验测试"),
    ], [1700, 4560, 3100])
    add_h2(doc, "7.2 后续实施里程碑")
    add_table(doc, ["阶段", "关键工作", "预期产出"], [
        ("产品稳定", "完善计划历史一致性、异常处理与演示模式", "可稳定演示的核心闭环"),
        ("用户验证", "招募首轮成年用户，观察理解度、行动意愿与焦虑反馈", "需求与可用性证据"),
        ("比赛打磨", "完成计划书、演示脚本、答辩材料和原创性说明", "完整参赛材料"),
        ("生态拓展", "研究更多设备、数据平台与场景的接入路径", "跨设备平台化路线图"),
    ], [1600, 4600, 3160])

    add_h1(doc, "08  团队建设与资源需求")
    add_h2(doc, "8.1 团队分工")
    add_table(doc, ["角色", "主要职责", "团队成员"], [
        ("项目负责人", "产品定位、整体推进、答辩与资源协调", "【待填写】"),
        ("产品与设计", "用户研究、交互设计、视觉与内容表达", "【待填写】"),
        ("技术研发", "健康数据、算法、AI、客户端与服务端开发", "【待填写】"),
        ("市场与运营", "用户访谈、试点组织、竞品与推广策略", "【待填写】"),
        ("指导资源", "健康安全、创新创业、技术与赛事指导", "【待填写】"),
    ], [1700, 4880, 2780])
    add_h2(doc, "8.2 所需资源")
    add_para(doc, "项目后续需要用户测试招募、健康安全文案审阅、智能穿戴与健康数据平台适配测试、云端AI调用与安全服务、比赛展示设备及校内创新创业资源支持。团队将优先使用低成本、可验证的方式推进，不以未经验证的大规模投入替代产品验证。")

    add_h1(doc, "09  社会价值、风险与合规")
    add_h2(doc, "9.1 社会价值")
    add_table(doc, ["价值方向", "知衡的贡献"], [
        ("健康素养", "帮助普通用户理解个人健康数据与不确定性，而非被单项数字驱动焦虑。"),
        ("行为改善", "将大而空泛的健康目标拆解为低负担、可执行、可评估的微行动。"),
        ("可信AI", "以事实包、结构化输出与安全校验约束AI健康交互。"),
        ("健康公平", "探索面向校园、企业和社区的普惠健康促进服务。"),
    ], [2200, 7160])
    add_h2(doc, "9.2 主要风险与应对")
    add_table(doc, ["风险", "应对方案"], [
        ("数据不足或设备差异", "数据质量守门；数据不足时不输出强结论；后续分阶段适配更多设备与数据平台。"),
        ("AI编造或过度推断", "结构化事实包、输出校验、本地规则降级和固定安全测试集。"),
        ("医疗误导", "坚持非诊疗定位、低风险建议与紧急症状优先安全提示。"),
        ("隐私泄露", "本地优先、最小必要数据、明确授权与删除能力。"),
        ("项目范围过大", "聚焦“洞察到行为改变”闭环，按阶段验证后再扩展。"),
    ], [2450, 6910])
    add_callout(doc, "合规底线", "知衡的价值建立在可信与克制之上：数据不足时诚实说明，相关性不写成因果，不将AI回答包装为诊断或治疗。")

    add_h1(doc, "10  附录：提交前待补材料清单")
    add_para(doc, "以下材料应由团队在最终提交前补齐。没有证据的内容不要写成已完成成果。")
    checklist = [
        "学校名称、参赛赛道、负责人、团队成员、指导教师和联系方式。",
        "产品截图、演示视频链接或二维码，并确保全部使用演示数据或完成脱敏。",
        "用户访谈或测试记录，包括样本数、时间、方法和主要发现。",
        "竞品调研来源、行业数据来源和引用日期。",
        "软著、专利、论文、获奖、合作或媒体报道等证明材料，如确实拥有。",
        "预算、资金来源与财务预测的计算依据，如学校提交模板要求。",
        "开源参考、许可证与知衡原创部分说明。",
        "健康安全文案审阅记录，或明确标注尚未完成专业审阅。",
    ]
    for item in checklist:
        add_bullet(doc, "待补：" + item)
    add_h2(doc, "团队答辩可使用的一句话")
    add_callout(doc, "核心陈述", "知衡不是又一个展示健康数据的应用，而是一个把个人健康数据转化为可解释洞察、低风险行动和可验证个人方法的AI健康智能体。", color=BLUE)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)
    cover(doc)
    toc(doc)
    body(doc)
    doc.core_properties.title = "知衡国创赛项目计划书"
    doc.core_properties.subject = "从健康洞察到行为改变的AI个人健康智能体"
    doc.core_properties.author = "知衡项目团队"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
