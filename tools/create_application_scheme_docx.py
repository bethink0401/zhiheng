from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path("/Users/chenweiyang/Documents/Project")
ASSET = ROOT / "竞赛材料/应用方案图片"
OUTPUT = ROOT / "竞赛材料/知衡AI健康洞察与行动平台应用方案.docx"
ICON = ROOT / "Design/AppIcon-Zhiheng-v4-focus-liquid-glass-white.png"

NAVY = "18324A"
TEAL = "2A7A78"
ORANGE = "E8874A"
GRAY = "606A73"
LIGHT = "EEF3F5"
LINE = "C9D4D9"
WHITE = "FFFFFF"
BLACK = "111111"
FONT = "PingFang SC"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            edge_data = kwargs[edge]
            tag = "w:" + edge
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            for key in ("val", "sz", "space", "color"):
                if key in edge_data:
                    element.set(qn("w:" + key), str(edge_data[key]))


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_alt_text(shape, title, description):
    doc_pr = shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_run_font(run, name=FONT, size=10.5, bold=False, color=BLACK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)


def style_paragraph(paragraph, after=5, before=0, line=1.28, first_line=True):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if first_line:
        fmt.first_line_indent = Cm(0.74)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def add_body(doc, text, after=5, bold_prefix=None):
    p = doc.add_paragraph()
    style_paragraph(p, after=after)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_heading(doc, text, level=1, keep=True):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = keep
    p.paragraph_format.space_before = Pt(7 if level == 1 else 4)
    p.paragraph_format.space_after = Pt(5 if level == 1 else 3)
    r = p.add_run(text)
    set_run_font(r, size=17 if level == 1 else 12.5, bold=True, color=BLACK)
    return p


def add_kicker(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=9.5, bold=True, color=TEAL)
    return p


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, size=8.5, color=GRAY)
    return p


def add_figure(doc, path, width_cm, title, description, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0)
    shape = p.add_run().add_picture(str(path), width=Cm(width_cm))
    set_alt_text(shape, title, description)
    add_caption(doc, caption)


def add_table(doc, headers, rows, widths=None, font_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0].cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_repeat_table_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, WHITE)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=font_size, bold=True, color=NAVY)
        if widths:
            cell.width = Cm(widths[i])
    for ridx, row_data in enumerate(rows):
        row = table.add_row()
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=85, bottom=85)
            set_cell_shading(cell, WHITE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_run_font(r, size=font_size)
            if widths:
                cell.width = Cm(widths[i])
            empty = {"val": "nil", "sz": "0", "color": WHITE, "space": "0"}
            set_cell_border(cell, top=empty, bottom=empty, left=empty, right=empty)
    for cell in table.rows[0].cells:
        top = {"val": "single", "sz": "12", "color": NAVY, "space": "0"}
        bottom = {"val": "single", "sz": "7", "color": NAVY, "space": "0"}
        empty = {"val": "nil", "sz": "0", "color": WHITE, "space": "0"}
        set_cell_border(cell, top=top, bottom=bottom, left=empty, right=empty)
    for cell in table.rows[-1].cells:
        bottom = {"val": "single", "sz": "12", "color": NAVY, "space": "0"}
        empty = {"val": "nil", "sz": "0", "color": WHITE, "space": "0"}
        set_cell_border(cell, top=empty, bottom=bottom, left=empty, right=empty)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_numbered_process(doc, items):
    table = doc.add_table(rows=len(items), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, (title, desc) in enumerate(items, 1):
        ncell, tcell = table.rows[idx - 1].cells
        ncell.width = Cm(1.0)
        tcell.width = Cm(15.8)
        ncell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(ncell, top=70, bottom=70)
        set_cell_margins(tcell, top=70, bottom=70)
        set_cell_shading(ncell, TEAL)
        np = ncell.paragraphs[0]
        np.alignment = WD_ALIGN_PARAGRAPH.CENTER
        np.paragraph_format.space_after = Pt(0)
        nr = np.add_run(str(idx))
        set_run_font(nr, size=10, bold=True, color=WHITE)
        tp = tcell.paragraphs[0]
        tp.paragraph_format.space_after = Pt(0)
        tr = tp.add_run(title + "  ")
        set_run_font(tr, size=9.3, bold=True)
        dr = tp.add_run(desc)
        set_run_font(dr, size=9.3)
        empty = {"val": "nil", "sz": "0", "color": WHITE, "space": "0"}
        line = {"val": "single", "sz": "3", "color": LINE, "space": "0"}
        set_cell_border(ncell, top=empty, bottom=empty, left=empty, right=empty)
        set_cell_border(tcell, top=empty, bottom=line, left=empty, right=empty)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_screenshot_grid(doc, shots):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, (path, label, description) in enumerate(shots):
        cell = table.cell(0, i)
        cell.width = Cm(5.65)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_margins(cell, top=30, start=50, end=50, bottom=30)
        border = {"val": "nil", "sz": "0", "color": WHITE, "space": "0"}
        set_cell_border(cell, top=border, bottom=border, left=border, right=border)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        shape = p.add_run().add_picture(str(path), width=Cm(5.0))
        set_alt_text(shape, label, description)
        cp = cell.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(0)
        cr = cp.add_run(label)
        set_run_font(cr, size=9.2, bold=True, color=NAVY)
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.65)
section.bottom_margin = Cm(1.55)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)
section.header_distance = Cm(0.65)
section.footer_distance = Cm(0.65)
section.different_first_page_header_footer = True

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
normal.paragraph_format.line_spacing = 1.28
normal.paragraph_format.space_after = Pt(5)

for level, size in ((1, 17), (2, 12.5)):
    st = styles[f"Heading {level}"]
    st.font.name = FONT
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(BLACK)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    st._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)

title_style = styles["Title"]
title_style.font.name = FONT
title_style.font.size = Pt(27)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor.from_string(BLACK)
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
title_style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run("知衡 AI健康洞察与行动平台应用方案")
set_run_font(hr, size=8, color=GRAY)
footer = section.footer
fp = footer.paragraphs[0]
add_page_number(fp)
for run in fp.runs:
    set_run_font(run, size=8, color=GRAY)

# 封面
doc.add_paragraph().paragraph_format.space_after = Pt(35)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(22)
shape = p.add_run().add_picture(str(ICON), width=Cm(3.2))
set_alt_text(shape, "知衡应用图标", "知衡 AI 健康洞察与行动平台的应用图标")

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(12)
r = p.add_run("知衡 AI健康洞察与行动平台应用方案")
set_run_font(r, size=27, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(24)
r = p.add_run("让个人健康数据转化为可理解 可执行 可验证的日常行动")
set_run_font(r, size=13, color=TEAL)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(16)
r = p.add_run("应用方案")
set_run_font(r, size=15, bold=True, color=NAVY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.left_indent = Cm(2.2)
p.paragraph_format.right_indent = Cm(2.2)
p.paragraph_format.line_spacing = 1.5
r = p.add_run("知衡以 AI 健康协同引擎为核心，把授权健康数据、个人基线、主观感受和生活情境组织为可追溯事实，再形成解释、追问和低风险微计划。系统持续记录执行结果，帮助用户沉淀适合自己的有效方法。")
set_run_font(r, size=11.5, color=GRAY)

doc.add_paragraph().paragraph_format.space_after = Pt(45)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("2026")
set_run_font(r, size=10, color=GRAY)

add_page_break(doc)

# 1 2
add_kicker(doc, "WHY NOW")
add_heading(doc, "一 项目背景")
add_body(doc, "智能穿戴设备和健康平台已经能够持续记录睡眠、心率、活动、呼吸、体温及训练等信息，但多数产品仍停留在数字呈现和固定提醒阶段。用户看到越来越多指标，却很难判断某次变化是否值得关注，也难以把数据与疲劳、压力、加班、饮食或运动等真实生活情境联系起来。")
add_body(doc, "生成式 AI 为个人健康管理提供了新的交互方式。知衡首先由本地程序完成数据质量判断、个人基线计算和趋势识别，再由 AI 在受约束的健康事实包内解释变化、补充关键追问并生成可执行建议。AI 由此成为连接数据理解、行动制定和结果复盘的中枢，而不是脱离证据的通用问答入口。")
add_body(doc, "当前产品以 iPhone 为首个落地载体，读取用户授权的 Apple Health 和 Apple Watch 数据。整体架构面向更广泛的移动终端、可穿戴设备和健康数据平台设计，后续可在保持同一安全边界的前提下扩展 Android、国产穿戴设备及机构服务端。")

add_heading(doc, "二 当前痛点问题")
add_table(doc, ["编号", "痛点", "用户影响", "知衡应对"], [
    ("1", "数据多但缺少解释", "难以理解波动是否与个人状态有关", "AI 基于个人基线和时间范围生成可追溯解释"),
    ("2", "通用阈值忽略个体差异", "容易产生无效提醒或不必要焦虑", "优先比较用户与过去的自己，并设置质量门槛"),
    ("3", "客观数据与主观感受割裂", "数据正常时仍可能忽略真实不适", "把精力 压力 身体感受和生活事件纳入分析"),
    ("4", "建议泛化且难以执行", "用户知道问题却不知道下一步怎么做", "一次生成一个低风险短周期微计划"),
    ("5", "行动缺少结果验证", "无法知道建议是否真正适合自己", "结合完成率 主客观变化和数据质量形成复盘"),
], widths=[1.1, 3.4, 5.3, 6.7], font_size=8.6)

add_page_break(doc)

# 3 4
add_kicker(doc, "USER AND NEED")
add_heading(doc, "三 需求分析")
add_body(doc, "用户真正需要的不是更多数据，而是一个能够回答连续问题的智能系统：最近发生了什么，哪些变化值得继续观察，可能与哪些生活因素相关，现在最适合做什么，以及行动后是否出现了积极变化。知衡把这些问题组织为完整闭环，避免用户在多个记录工具和搜索入口之间反复切换。")
add_body(doc, "需求可归纳为四类。第一是可信理解，系统必须先判断授权、覆盖率、有效日和数据来源是否足以支持结论。第二是个体化解释，所有趋势优先基于个人基线，并同时尊重主观感受。第三是行动支持，建议需要低门槛、短周期、可暂停。第四是长期积累，系统要把多次执行结果沉淀为个人经验，而不是不断给出一次性建议。")
add_body(doc, "在医疗安全方面，用户需要明确知道系统能做什么和不能做什么。知衡定位为健康管理和生活方式改善工具，不提供疾病诊断、治疗方案或药物剂量调整。出现紧急症状时，独立安全规则优先提示用户寻求当地急救或专业医疗帮助。")

add_heading(doc, "四 目标用户群体")
add_table(doc, ["用户群体", "典型场景", "核心需求", "产品价值"], [
    ("关注状态变化的成年人", "睡眠波动 精力下降 压力增加", "理解变化并获得轻量建议", "把分散数据转化为可执行行动"),
    ("已有穿戴设备的用户", "长期积累数据但缺少复盘", "看懂个人趋势和影响因素", "建立个人基线与长期方法库"),
    ("生活节奏不稳定的人群", "加班 出差 聚餐 训练变化", "连接生活事件与身体状态", "通过背景记录减少机械判断"),
    ("健康管理服务提供方", "需要低风险数字化随访工具", "提升陪伴效率并保持安全边界", "形成可配置的机构服务能力"),
], widths=[3.6, 4.1, 4.4, 4.5], font_size=8.8)
add_body(doc, "产品以成年人日常健康管理为主要场景。对于老年人、慢病人群或康复人群，知衡可在专业机构指导下提供记录、提醒和趋势整理，但不会替代临床判断。")

add_page_break(doc)

# 5
add_kicker(doc, "PRODUCT SCOPE")
add_heading(doc, "五 功能需求")
add_body(doc, "功能设计围绕一个连续闭环展开，并把 AI 放在理解和协同的核心位置。")
add_table(doc, ["核心能力", "功能说明", "AI 作用", "输出结果"], [
    ("健康事实理解", "读取授权数据并计算质量 个人基线和趋势", "将结构化事实转化为自然语言解释并标注不确定性", "今日状态 趋势卡 依据范围"),
    ("主观与情境补充", "记录精力 压力 身体感受和生活事件", "识别解释缺口并最多提出一个关键追问", "更贴近真实生活的健康上下文"),
    ("微计划生成与执行", "提供三至七天的低风险行动并记录完成情况", "结合事实包筛选适合当前状态的行动方向", "单一微计划 每日任务 反馈记录"),
    ("结果评估与方法沉淀", "比较计划前后变化 完成率和数据质量", "对确定性结果进行易懂解读并提示限制", "计划评估 我的有效方法"),
    ("隐私与安全控制", "权限管理 本地存储 数据删除和最小化外发", "仅消费白名单聚合事实并接受输出校验", "可追溯来源 安全降级 用户控制"),
], widths=[3.0, 5.2, 5.2, 3.2], font_size=8.6)
add_body(doc, "用户可以从今日状态进入洞察，也可以直接向 AI 助手提问。无论入口如何变化，系统都使用同一套数据质量、趋势和安全边界，避免页面之间出现互相矛盾的结论。")

add_page_break(doc)

# 6 7 architecture
add_kicker(doc, "ENGINEERING")
add_heading(doc, "六 开发工具与技术环境")
add_body(doc, "当前版本使用 Xcode 和 Swift 进行开发，以 SwiftUI 构建界面，HealthKit 读取用户授权的健康数据，CareKitStore 管理微计划、任务日程和执行结果，SwiftData 保存主观感受、生活事件、洞察快照和个人方法，Swift Charts 展示趋势，UserNotifications 提供低敏感提醒。网络层采用 URLSession，AI 服务通过可替换协议接入，敏感配置与业务代码分离。")
add_body(doc, "开发过程采用真实数据提供者、隔离示例数据提供者和测试数据提供者共用同一协议的方式，使算法、页面和安全规则可以独立验证。当前工程已形成移动端闭环，后续可将设备适配层扩展至其他可穿戴平台，并通过服务端协同支持跨终端账户、机构版管理和模型路由。")

add_heading(doc, "七 技术方案")
add_figure(doc, ASSET / "png/02-system-architecture.png", 16.7, "知衡系统技术架构", "从设备和用户输入到本地数据层 分析引擎 AI 协同层 行动闭环与应用层的系统架构", "图 1  知衡系统技术架构")
add_body(doc, "架构遵循本地优先和最小必要原则。原始健康样本留在设备和系统健康库中，本地分析引擎先完成清洗、有效日统计、二十八天基线、七天窗口、稳健变化和数据质量分级。只有用户主动发起 AI 请求时，系统才发送回答所需的聚合事实，不包含底层标识、精确事件时间或完整原始样本。")

add_page_break(doc)

add_kicker(doc, "TECHNICAL FLOW")
add_body(doc, "本地确定性分析。知衡使用中位数和稳健变化方法降低单日异常值的影响。最近七天至少四个有效日才形成短期趋势，二十八天至少十四个有效日才称为个人基线。缺失数据不按零处理，未佩戴、权限变化和数据源变化也不会自动解释为健康恶化。固定输入产生固定结果，AI 不替代本地程序计算趋势。", bold_prefix="本地确定性分析。")
add_body(doc, "受约束的 AI 协同。AI 服务接收结构化健康事实包，事实包包括时间范围、数据质量、当前值、基线、程序计算的趋势、用户主动提供的主观感受和生活情境，以及允许推荐的低风险计划模板。模型输出经过客户端规则校验，区分事实、可能因素、不确定性、追问、建议行动和安全级别。网络异常、超时或格式错误时，应用自动降级为本地规则摘要。", bold_prefix="受约束的 AI 协同。")
add_body(doc, "行动与验证。微计划由 CareKitStore 作为唯一任务事实来源，记录计划、日程和每日结果；SwiftData 保存计划开始时的聚合基线、同期生活背景和最终分析快照。评估同时查看客观指标、主观感受、完成率和数据质量，只表达可能有帮助、暂未观察到明显变化、执行不足或数据不足，避免把相关性写成因果关系。", bold_prefix="行动与验证。")
add_body(doc, "隐私与安全。用户可以查看权限状态、删除主观记录、生活事件、计划数据和 AI 对话。提醒默认关闭，锁屏只显示通用标题。系统不使用健康数据投放广告或出售画像，也不把健康数据用于训练基础模型。面向其他终端扩展时，这些边界保持不变。", bold_prefix="隐私与安全。")
add_table(doc, ["数据或能力", "主要位置", "处理规则"], [
    ("原始健康样本", "HealthKit 或设备健康平台", "只读授权 默认不上传"),
    ("趋势和数据质量", "本地分析引擎", "确定性计算 可测试 可追溯"),
    ("AI 健康事实包", "加密网络请求", "用户主动触发 最小字段 白名单限制"),
    ("计划和执行结果", "CareKitStore", "单一事实来源 可暂停 可删除"),
    ("感受 事件 方法", "本地持久化", "用户控制 支持删除和隔离示例数据"),
], widths=[4.0, 4.2, 8.2], font_size=8.9)

add_page_break(doc)

# 8 screenshots page 1
add_kicker(doc, "CORE EXPERIENCE")
add_heading(doc, "八 作品功能说明")
add_body(doc, "产品把复杂分析放在后台，把前台体验收敛为看状态、读洞察、问 AI、做计划和留方法五个动作。每个页面都能展开查看所用数据范围和质量，用户不需要理解统计模型，也能知道结论从哪里来。")
add_screenshot_grid(doc, [
    (ASSET / "产品截图/01-today.png", "今日状态", "今日页面展示个人状态摘要 健康指标和主观签到入口"),
    (ASSET / "产品截图/02-insights.png", "可解释洞察", "洞察页面展示数据事实 可能解释 建议和依据"),
    (ASSET / "产品截图/03-ai-assistant.png", "AI 健康助手", "AI 助手基于聚合健康事实进行流式问答并展示依据"),
])
add_body(doc, "今日状态与洞察。今日页汇总近期变化、数据质量和主观签到入口。洞察页把结论分成事实、可能因素、建议行动和依据来源；当数据不足时，页面直接说明原因并停止输出强结论。用户可以从洞察进入详情，也可以把当前问题带入 AI 对话。", bold_prefix="今日状态与洞察。")
add_body(doc, "AI 健康助手。用户可用自然语言询问睡眠、活动、精力、压力和近期计划等问题。助手读取限定的健康事实包，生成带时间范围和不确定性说明的回答；当事实不足时优先追问一个最关键背景。对于诊断、药物或紧急症状请求，安全规则限制回答范围并提供就医提示。", bold_prefix="AI 健康助手。")

add_page_break(doc)

# 8 screenshots page 2
add_kicker(doc, "ACTION AND LEARNING")
add_screenshot_grid(doc, [
    (ASSET / "产品截图/04-micro-plan-active.png", "进行中微计划", "微计划页面展示短周期任务 日程 完成和跳过状态"),
    (ASSET / "产品截图/05-plan-evaluation.png", "计划结束评估", "评估页面比较执行情况 客观趋势 主观变化和数据质量"),
    (ASSET / "产品截图/06-effective-methods.png", "我的有效方法", "方法页面按证据维度沉淀用户历史计划和适用方法"),
])
add_body(doc, "低风险微计划。系统一次只建议一个三至七天行动，例如固定上床准备时间、午后短时步行或训练恢复安排。用户确认后才创建计划，可每日完成或跳过并留下可选反馈，也可以暂停、恢复或提前结束。计划建议控制在生活方式改善范围内，不涉及诊断、处方和药物调整。", bold_prefix="低风险微计划。")
add_body(doc, "计划评估与我的有效方法。计划结束后，系统根据真实执行记录生成评估，分别展示完成率、客观变化、主观变化和数据质量。多次结果会沉淀为方法卡，标明证据维度与可信度。用户可以隐藏、恢复或重新验证某个方法，新计划会建立独立基线，不改写旧历史。", bold_prefix="计划评估与我的有效方法。")
add_body(doc, "应用还提供历史时间线、智能提醒、匿名健康报告和隐私与安全中心。对于尚未连接健康设备或暂时没有足够数据的用户，系统可以提供明确标注的示例内容，帮助用户了解主要功能。示例内容与真实个人数据完全隔离，通知不会显示敏感健康信息，报告只有在用户确认后才能通过系统功能分享。")

add_page_break(doc)

# 9 usage
add_kicker(doc, "USER JOURNEY")
add_heading(doc, "九 使用说明与操作流程")
add_figure(doc, ASSET / "png/03-user-journey.png", 16.7, "知衡用户操作流程", "用户从授权与签到到查看洞察 咨询 AI 确认计划 执行评估和沉淀方法的操作流程", "图 2  用户操作流程")
add_numbered_process(doc, [
    ("了解产品并完成授权", "首次进入后阅读产品定位、AI 工作方式和隐私边界，再按需授权健康数据。拒绝部分权限不会阻止使用其他功能。"),
    ("补充主观状态", "在今日页记录精力 压力和身体感受，也可添加加班 出差 饮食或训练等生活事件。"),
    ("查看今日与洞察", "阅读近期趋势和数据质量，进入详情查看基线 当前窗口 样本天数和依据。"),
    ("向 AI 提问", "输入自然语言问题。系统显示正在使用的事实类别和时间范围，并以流式方式返回回答。"),
    ("确认微计划", "查看行动内容 周期和安全边界。只有用户确认后才创建计划。"),
    ("每日执行", "在计划页选择完成或跳过，可补充简短反馈，也可随时暂停或提前结束。"),
    ("查看评估", "计划结束后阅读客观 主观 执行和数据质量四类结果，决定是否继续观察或重新验证。"),
    ("管理数据", "在设置中查看权限 隐私说明 通知和导出选项，按需删除本地记录与对话。"),
])
add_body(doc, "交互文案采用低焦虑表达，不使用红色恐吓、综合健康分或确定疾病措辞。所有关键操作都保留用户确认，AI 不会自动修改计划、发送敏感通知或替代用户作出医疗决定。")

add_page_break(doc)

# 10 AI
add_kicker(doc, "AI FIRST")
add_heading(doc, "十 AI的核心作用与项目创新")
add_figure(doc, ASSET / "png/01-ai-health-closed-loop.png", 16.7, "AI 健康行动闭环", "知衡将多源健康信息经过本地质量守门和个人基线分析后交给 AI 协同引擎 再形成微计划 结果评估和个人方法", "图 3  AI 驱动的健康行动闭环")
add_body(doc, "知衡的 AI 不是附加聊天功能，而是贯穿理解、交互、行动和复盘的协同引擎。它首先把本地算法形成的结构化事实转换为用户能理解的语言，再根据解释缺口提出一个关键追问，并从低风险模板中组织最适合当前情境的微计划。计划结束后，AI 负责解释程序已经计算出的评估结果，帮助用户理解哪些变化值得继续观察。")
add_body(doc, "项目的第一项创新是个人基线优先。系统主要比较用户与过去的自己，而不是用单一通用阈值评价所有人。第二项创新是数据质量守门，覆盖率和来源不足时，AI 必须承认不确定性。第三项创新是主客观与情境共同进入事实包，使数据和真实感受具有同等地位。第四项创新是从建议走向验证，用户每次行动都会留下可追溯结果，并逐步形成个人方法库。")
add_body(doc, "面向后续发展，AI 引擎可接入多模型路由、语音交互和更多设备数据，并通过长期记忆形成更连续的健康陪伴。但扩展仍以用户授权、最小必要数据和专业安全规则为前提。当前版本不会声称疾病诊断或持续急救监护能力。")

add_page_break(doc)

# 11 12
add_kicker(doc, "GROWTH")
add_heading(doc, "十一 应用前景")
add_body(doc, "个人健康数据正在从偶发测量转向连续记录，但数据解释与行动服务仍存在明显空缺。知衡可覆盖日常状态管理、运动恢复、睡眠习惯改善、职场健康和轻量康复记录等场景。随着设备适配层扩展，产品可以服务更多手机和可穿戴设备用户，并保持统一的个人基线、AI 事实包和行动验证框架。")
add_body(doc, "在个人市场，知衡可以成为用户长期使用的健康入口；在机构市场，可为高校、企业健康项目、保险健康服务、运动机构和基层健康管理提供可配置能力。机构侧获得的是经过授权的趋势摘要和服务工具，而不是未经同意的原始个人数据。")

add_heading(doc, "十二 商业模式")
add_table(doc, ["模式", "服务对象", "主要内容", "收入方式"], [
    ("个人基础服务", "普通用户", "健康数据汇总 主观记录 基础趋势和本地微计划", "免费使用 扩大用户基础"),
    ("个人订阅服务", "有持续管理需求的用户", "高级 AI 解读 长周期复盘 跨设备同步和个性化报告", "月度或年度订阅"),
    ("机构服务", "企业 高校 运动与健康机构", "人群健康项目配置 匿名化汇总和服务流程管理", "按账户或服务周期收费"),
    ("平台能力输出", "设备厂商与数字健康合作方", "数据适配 AI 事实包 行动计划和评估组件", "授权费 接口费或联合运营分成"),
], widths=[3.0, 3.8, 6.5, 3.3], font_size=8.8)
add_body(doc, "商业化坚持三条边界：健康数据不用于广告定向，不出售个人画像，不以制造焦虑提高付费转化。高级服务的价值来自更连续的解释、更多设备连接和更系统的复盘，而不是隐藏基础安全信息。")

add_page_break(doc)

# 13
add_kicker(doc, "POSITIONING")
add_heading(doc, "十三 市场竞争与项目优势")
add_body(doc, "市场中的健康应用大致分为数据展示、通用 AI 问答和习惯打卡三类。知衡把三类能力连接为同一条可验证路径，并通过个人基线、数据质量和安全规则限制 AI 的推断范围。")
add_table(doc, ["比较维度", "数据展示类应用", "通用 AI 助手", "习惯打卡类应用", "知衡"], [
    ("数据来源", "设备或平台数据", "主要依赖用户描述", "主要依赖手动记录", "授权健康数据与主观情境结合"),
    ("个体化依据", "常用固定阈值", "依赖对话提示", "依赖预设目标", "二十八天个人基线与七天当前窗口"),
    ("AI 可追溯性", "通常无 AI", "依据可能不透明", "通常无 AI", "展示事实类别 时间范围和不确定性"),
    ("行动能力", "提醒或报告", "生成泛化建议", "记录既定习惯", "一次一个低风险短周期微计划"),
    ("效果验证", "少量趋势回顾", "通常不连续跟踪", "关注打卡天数", "客观 主观 完成率和质量共同评估"),
    ("长期资产", "历史数据", "对话记录", "连续打卡记录", "可重新验证的个人有效方法"),
    ("安全边界", "依产品而异", "需要额外约束", "风险相对较低", "本地质量守门 输出校验和紧急规则"),
], widths=[2.4, 3.35, 3.35, 3.35, 4.35], font_size=8.1)
add_body(doc, "项目优势来自闭环而不是单个页面。健康数据先被验证，再被解释；建议由用户确认后执行；效果由确定性程序评估；AI 只在有依据的范围内帮助用户理解。这一结构也便于扩展到其他设备和服务场景，因为数据接入层可以变化，核心事实模型和安全规则保持稳定。")

add_page_break(doc)

# 14
add_kicker(doc, "DELIVERY")
add_heading(doc, "十四 项目可行性与落地计划")
add_body(doc, "项目已经完成 iPhone 端核心闭环，包括健康数据授权与趋势、主观记录、生活事件、可解释洞察、真实流式 AI 对话、CareKit 微计划、结束评估、我的有效方法、通知和隐私说明。数据层、算法层、AI 层和界面层通过协议隔离，便于测试和替换。现有实现为进一步用户验证、设备扩展和机构合作提供了可运行基础。")
add_body(doc, "技术可行性来自成熟平台能力与明确分工：HealthKit 提供授权健康数据，CareKitStore 负责行动记录，SwiftData 管理本地上下文，确定性算法保证趋势可复现，AI 负责解释与交互。运营可行性来自低门槛微计划和本地优先隐私策略，用户无需购买新硬件即可从已有数据开始。")

add_page_break(doc)

# 15 risks
add_kicker(doc, "RISK CONTROL")
add_heading(doc, "十五 风险分析与应对措施")
add_table(doc, ["风险", "可能影响", "应对措施"], [
    ("健康数据不足或来源变化", "趋势失真或引发误解", "设置有效日和覆盖率门槛 识别权限与来源变化 数据不足时停止强结论"),
    ("AI 产生越界或不准确回答", "误导用户或形成医疗风险", "限定结构化事实包 使用提示注入防护 输出校验 本地降级和独立紧急规则"),
    ("隐私泄露", "损害用户权益和产品信任", "本地优先 最小化外发 加密传输 权限说明 删除能力和敏感日志审计"),
    ("用户难以长期坚持", "计划完成率低 方法积累不足", "一次只执行一个短计划 降低操作成本 允许暂停 跳过和反馈"),
    ("设备与平台差异", "跨终端数据口径不一致", "建立统一领域模型和设备适配层 对每类指标记录来源与质量"),
    ("商业化损害产品中立性", "诱导消费或健康焦虑", "禁止健康广告画像 基础安全信息不设付费墙 建立合作方准入规则"),
    ("专业适用范围不清", "用户把产品当作诊疗工具", "持续显示非诊断定位 高风险场景转介 专业医疗文案审阅"),
], widths=[4.0, 5.0, 7.8], font_size=8.7)
add_body(doc, "风险控制贯穿产品流程，而不是依赖免责声明。数据质量在分析前执行，安全规则在模型调用前后执行，用户确认在计划创建前执行，删除和权限控制在设置中持续可见。后续每新增一种设备、指标或合作场景，都需要重新完成数据、隐私和医疗安全评估。")

add_page_break(doc)

# 16 17 18 19
add_kicker(doc, "IMPACT")
add_heading(doc, "十六 社会价值")
add_body(doc, "知衡帮助普通用户建立更理性的健康数据观。系统承认数据缺失和不确定性，减少因单次波动产生的焦虑，也避免用设备数据否定真实感受。通过低风险行动和结果复盘，用户可以逐步形成更可持续的睡眠、活动和恢复习惯。")
add_body(doc, "对于高校、企业和基层健康服务，知衡提供可配置的数字化陪伴框架，使有限的专业资源更集中地用于需要人工关注的人群。对于数字健康行业，项目提供一种可复用路径：让 AI 使用结构化、最小化、可追溯的健康事实，并把解释连接到用户确认和长期验证。")

add_heading(doc, "十七 产品发展规划")
add_table(doc, ["阶段", "重点目标", "主要交付"], [
    ("第一阶段", "完善核心体验和小规模用户验证", "稳定移动端闭环 优化解释质量 完成医疗文案审阅"),
    ("第二阶段", "扩展数据来源和交互方式", "适配更多穿戴设备 支持语音和跨设备同步"),
    ("第三阶段", "形成机构服务能力", "企业与高校健康项目后台 匿名化汇总和配置工具"),
    ("第四阶段", "建设开放健康行动平台", "标准化设备接入 模型路由和合作方组件"),
], widths=[2.6, 6.0, 8.2], font_size=8.9)
add_body(doc, "各阶段以安全门禁和验证证据作为放行条件。跨平台扩展不会简单复制界面，而是先统一数据语义、授权状态、质量计算和删除机制，再开放新的终端入口。")

add_heading(doc, "十八 预期成果")
add_body(doc, "近期成果是完成稳定可用的个人健康闭环，持续提高数据覆盖、洞察可理解性、微计划完成率和结果复盘质量。中期成果是支持更多设备与终端，形成跨平台个人健康事实模型和机构服务能力。长期成果是沉淀经过用户反复验证的个人方法库，让健康建议从一次性内容转变为可追踪、可修正的个人经验。")

add_heading(doc, "十九 项目总结")
add_body(doc, "知衡以 AI 健康协同引擎为核心，把健康数据、个人基线、主观感受和生活情境连接为一套可理解的事实体系，再通过低风险微计划把解释转化为行动，通过客观与主观结果共同评估行动效果。系统既发挥 AI 的语言理解和个性化交互能力，也用本地确定性算法、数据质量门槛、用户确认和安全规则约束其边界。")
add_body(doc, "当前移动端产品已经形成从数据到方法的完整路径，技术架构也为更多手机、穿戴设备和机构服务预留了扩展空间。知衡希望让每个人都能在尊重隐私和不确定性的前提下，更清楚地理解自己的状态，更有把握地采取下一步行动，并把一次次实践积累成真正适合自己的健康方法。")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(16)
r = p.add_run("知衡 让健康数据成为行动依据")
set_run_font(r, size=13, bold=True, color=TEAL)

doc.core_properties.title = "知衡 AI健康洞察与行动平台应用方案"
doc.core_properties.subject = "AI健康洞察 行动计划与结果验证应用方案"
doc.core_properties.keywords = "知衡 AI 健康管理 个人基线 微计划"
doc.core_properties.creator = ""
doc.core_properties.last_modified_by = ""
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
